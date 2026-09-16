import openpyxl
import re

# 1. Map all real content URLs from 브랜드슬램 x OWM.xlsx
wb_owm = openpyxl.load_workbook("data/excel/브랜드슬램 x OWM.xlsx", data_only=True)

# Sheet: 6월 영미권 OWM(강남, 북촌, 성수, 종각, 이태원)
ws_june_en = wb_owm["6월 영미권 OWM(강남, 북촌, 성수, 종각, 이태원)"]
headers = [ws_june_en.cell(row=1, column=c).value for c in range(1, ws_june_en.max_column+1)]

name_idx = headers.index("name") + 1
loc_idx = 1
insta_prof_idx = [i for i, h in enumerate(headers) if h == "insta"][0] + 1
tt_prof_idx = [i for i, h in enumerate(headers) if h == "tiktok"][0] + 1
tt_post_idx = [i for i, h in enumerate(headers) if h and "Posting URL (TT)" in str(h)][0] + 1
ig_post_idx = [i for i, h in enumerate(headers) if h and "Posting URL (IG)" in str(h)][0] + 1

june_en_data = []
for r in range(2, ws_june_en.max_row+1):
    loc_val = ws_june_en.cell(row=r, column=loc_idx).value
    name_val = ws_june_en.cell(row=r, column=name_idx).value
    if not name_val:
        continue
    name = str(name_val).strip()
    loc = str(loc_val).strip() if loc_val else ""
    insta_prof = ws_june_en.cell(row=r, column=insta_prof_idx).value
    tt_prof = ws_june_en.cell(row=r, column=tt_prof_idx).value
    tt_post = ws_june_en.cell(row=r, column=tt_post_idx).value
    ig_post = ws_june_en.cell(row=r, column=ig_post_idx).value
    
    june_en_data.append({
        "row": r,
        "loc": loc,
        "name": name,
        "insta_prof": str(insta_prof).strip() if insta_prof else None,
        "tt_prof": str(tt_prof).strip() if tt_prof else None,
        "tt_post": str(tt_post).strip() if tt_post else None,
        "ig_post": str(ig_post).strip() if ig_post else None,
    })

print(f"Loaded {len(june_en_data)} influencers from 6월 영미권 sheet.")
for item in june_en_data:
    print(f"{item['name']} | IG post: {item['ig_post']} | TT post: {item['tt_post']}")

