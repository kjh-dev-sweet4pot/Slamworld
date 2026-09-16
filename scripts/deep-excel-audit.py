import openpyxl, os, re

def norm(s):
    if s is None:
        return ""
    return str(s).strip()

def clean_url(u):
    if not u:
        return ""
    return str(u).strip()

print("=========================================================")
print("1. Auditing 브랜드슬램 x OWM.xlsx")
print("=========================================================")
wb1 = openpyxl.load_workbook("data/excel/브랜드슬램 x OWM.xlsx", data_only=True)

for sheetname in wb1.sheetnames:
    ws = wb1[sheetname]
    if ws.max_row < 2:
        continue
    # look for posting url or url columns
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column+1)]
    if not any(headers):
        headers = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column+1)]
    
    url_cols = [(i+1, h) for i, h in enumerate(headers) if h and ("url" in str(h).lower() or "posting" in str(h).lower())]
    if url_cols:
        print(f"\nSheet [{sheetname}] (rows: {ws.max_row})")
        print(f"  URL Columns: {url_cols}")

print("\n=========================================================")
print("2. Auditing OWM 명동 - 본시트.xlsx")
print("=========================================================")
wb2 = openpyxl.load_workbook("data/excel/OWM 명동 - 본시트.xlsx", data_only=True)
for sheetname in wb2.sheetnames:
    ws = wb2[sheetname]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column+1)]
    url_cols = [(i+1, h) for i, h in enumerate(headers) if h and ("url" in str(h).lower() or "posting" in str(h).lower())]
    print(f"\nSheet [{sheetname}] (rows: {ws.max_row})")
    print(f"  Headers: {[h for h in headers if h is not None][:12]}")
    print(f"  URL Columns: {url_cols}")

